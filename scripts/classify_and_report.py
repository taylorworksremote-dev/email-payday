import argparse, json, sys, re, html as _html
from pathlib import Path
from datetime import datetime, timedelta, timezone
from email.utils import parsedate_to_datetime
try:
    from zoneinfo import ZoneInfo
except ImportError:
    from backports.zoneinfo import ZoneInfo

KEYWORDS = [
    "payment due",
    "late payment",
    "payment reminder",
    "due",
    "make a payment",
    "past due",
    "pay now",
]

_KEYWORD_PATTERNS = [(kw, re.compile(re.escape(kw), re.IGNORECASE)) for kw in KEYWORDS]


def find_matches(text):
    if not text:
        return []
    return [kw for kw, pat in _KEYWORD_PATTERNS if pat.search(text)]


def parse_date_any(date_raw):
    """Best-effort parse of a date string from either the Gmail API or an RFC 2822
    email header into an aware UTC datetime. Returns None if unparseable."""
    if date_raw is None:
        return None
    if isinstance(date_raw, (int, float)):
        ts = date_raw / 1000 if date_raw > 10**12 else date_raw
        return datetime.fromtimestamp(ts, tz=timezone.utc)
    s = str(date_raw).strip()
    if not s:
        return None
    try:
        dt = parsedate_to_datetime(s)
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=timezone.utc)
        return dt.astimezone(timezone.utc)
    except Exception:
        pass
    try:
        dt = datetime.fromisoformat(s.replace("Z", "+00:00"))
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=timezone.utc)
        return dt.astimezone(timezone.utc)
    except Exception:
        pass
    try:
        ts = float(s)
        ts = ts / 1000 if ts > 10**12 else ts
        return datetime.fromtimestamp(ts, tz=timezone.utc)
    except Exception:
        pass
    return None


def guess_company(sender_email, sender_name):
    domain = (sender_email or "").split("@")[-1].lower()
    generic = {"gmail.com", "yahoo.com", "icloud.com", "outlook.com", "hotmail.com",
               "aol.com", "me.com", "protonmail.com"}
    if domain and domain not in generic:
        parts = domain.split(".")
        core = parts[-2] if len(parts) >= 2 else parts[0]
        return core.replace("-", " ").title()
    return (sender_name or "").strip() or "Unknown"


def load_records(paths):
    records = []
    for p in paths:
        p = Path(p)
        files = list(p.glob("*.json")) if p.is_dir() else [p]
        for f in files:
            try:
                data = json.loads(f.read_text())
                if isinstance(data, list):
                    records.extend(data)
            except Exception as e:
                print(f"Skipping {f}: {e}", file=sys.stderr)
    return records


def render_email_screenshot(subject, sender_name, sender_email, date_local_str, body_html, body_text, out_path):
    """Render a visual copy of the email (header block + body) to a PNG using a headless
    browser, so the report can carry actual evidence of the message rather than just text --
    useful when the sheet is going to be handed to an attorney. Uses the real HTML body when
    the collector captured one; otherwise falls back to a formatted view of the plain text.
    Returns out_path on success, or None if Playwright isn't installed or rendering fails
    (the rest of the report still gets built either way -- this is best-effort)."""
    try:
        from playwright.sync_api import sync_playwright
    except ImportError:
        return None

    if body_html and body_html.strip():
        body_section = body_html
    else:
        body_section = f"<pre style=\"white-space:pre-wrap;font-family:inherit;margin:0;\">{_html.escape(body_text or '')}</pre>"

    page_html = f"""<!doctype html>
<html><head><meta charset="utf-8"></head>
<body style="margin:0;background:#f2f2f2;">
  <div style="max-width:760px;margin:0 auto;padding:20px;background:#fff;
              font-family:-apple-system,'Segoe UI',Arial,sans-serif;">
    <div style="font-size:16px;font-weight:600;margin-bottom:10px;color:#111;">
      {_html.escape(subject) or "(no subject)"}
    </div>
    <div style="font-size:13px;color:#444;margin-bottom:2px;">
      <b>From:</b> {_html.escape(sender_name or "")} &lt;{_html.escape(sender_email or "")}&gt;
    </div>
    <div style="font-size:13px;color:#444;margin-bottom:14px;">
      <b>Date:</b> {_html.escape(date_local_str or "")}
    </div>
    <hr style="border:none;border-top:1px solid #eee;margin-bottom:14px;">
    <div style="font-size:14px;line-height:1.5;color:#222;">{body_section}</div>
  </div>
</body></html>"""

    try:
        with sync_playwright() as p:
            browser = p.chromium.launch()
            page = browser.new_page(viewport={"width": 800, "height": 100})
            page.set_content(page_html, wait_until="networkidle", timeout=15000)
            page.screenshot(path=str(out_path), full_page=True)
            browser.close()
        return out_path
    except Exception as e:
        print(f"  Screenshot failed for '{subject}': {e}", file=sys.stderr)
        return None


def write_outputs(rows, out_prefix):
    import csv
    columns = ["date_local_str", "day_of_week", "account", "folder", "sender_name",
               "sender_email", "company_guess", "subject", "matched_keywords",
               "potential_fdcpa_flag", "excerpt", "source_ref", "screenshot_path"]
    csv_path = f"{out_prefix}.csv"
    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=columns)
        writer.writeheader()
        for r in rows:
            writer.writerow({c: r.get(c, "") for c in columns})
    print(f"Wrote {csv_path}", file=sys.stderr)

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.drawing.image import Image as XLImage
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("openpyxl not installed (pip install openpyxl --break-system-packages) "
              "-- CSV was still written.", file=sys.stderr)
        return

    headers = ["Date/Time (Local)", "Day", "Account", "Folder", "Sender Name", "Sender Email",
               "Company (guess)", "Subject", "Matched Keyword(s)", "Potential FDCPA Flag",
               "Excerpt", "Reference", "Email Screenshot"]
    screenshot_col = len(headers)  # 1-indexed column number for the image column
    screenshot_col_letter = get_column_letter(screenshot_col)

    wb = Workbook()
    ws_all = wb.active
    ws_all.title = "All Matches"
    ws_flag = wb.create_sheet("Potential FDCPA Claims")
    fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    bold = Font(bold=True)
    for ws in (ws_all, ws_flag):
        ws.append(headers)
        for cell in ws[1]:
            cell.font = bold
        ws.column_dimensions[screenshot_col_letter].width = 46

    def append_row(ws, r, vals):
        ws.append(vals)
        row_num = ws.max_row
        path = r.get("screenshot_path")
        if path and Path(path).exists():
            try:
                img = XLImage(path)
                # Scale down to a fixed width so the sheet stays a manageable size.
                target_w = 320
                if img.width:
                    scale = target_w / img.width
                    img.width = target_w
                    img.height = int(img.height * scale)
                ws.row_dimensions[row_num].height = max(img.height * 0.75, 15)
                ws.add_image(img, f"{screenshot_col_letter}{row_num}")
            except Exception as e:
                print(f"  Could not embed screenshot for row {row_num}: {e}", file=sys.stderr)
        return row_num

    for r in rows:
        vals = [r["date_local_str"], r["day_of_week"], r["account"], r["folder"],
                r["sender_name"], r["sender_email"], r["company_guess"], r["subject"],
                r["matched_keywords"], r["potential_fdcpa_flag"], r["excerpt"], r["source_ref"],
                "" if r.get("screenshot_path") else "(no screenshot)"]
        row_num = append_row(ws_all, r, vals)
        if r["potential_fdcpa_flag"] == "Yes":
            for cell in ws_all[row_num]:
                cell.fill = fill
            append_row(ws_flag, r, vals)

    for ws in (ws_all, ws_flag):
        for col_cells in ws.columns:
            col_letter = col_cells[0].column_letter
            if col_letter == screenshot_col_letter:
                continue
            length = max((len(str(c.value)) if c.value else 0) for c in col_cells)
            ws.column_dimensions[col_letter].width = min(max(length + 2, 10), 60)

    xlsx_path = f"{out_prefix}.xlsx"
    wb.save(xlsx_path)
    print(f"Wrote {xlsx_path}", file=sys.stderr)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", nargs="+", required=True,
                     help="One or more raw JSON files (or directories of them)")
    ap.add_argument("--out-prefix", required=True)
    ap.add_argument("--years", type=float, default=2.0)
    ap.add_argument("--tz", required=True,
                     help="IANA timezone for the user's actual location, e.g. America/New_York, "
                          "America/Chicago, America/Denver, America/Los_Angeles, America/Phoenix. "
                          "Required -- don't default this to Eastern time; ask the user where "
                          "they live and map it to the right zone (their state may span more "
                          "than one, e.g. Texas or Indiana, so confirm the city/region if so).")
    ap.add_argument("--flag-start-hour", type=int, default=21)
    ap.add_argument("--flag-end-hour", type=int, default=8)
    ap.add_argument("--screenshots", choices=["all", "flagged", "none"], default="all",
                     help="Which rows get an embedded screenshot of the actual email. "
                          "'all' (default) captures every match; 'flagged' only the "
                          "potential-FDCPA (off-hours) rows, which is faster if you only "
                          "need evidence for those; 'none' skips screenshots entirely.")
    ap.add_argument("--screenshot-dir", default=None,
                     help="Where to save screenshot PNGs before they're embedded in the xlsx. "
                          "Defaults to '<out-prefix>_screenshots/'.")
    args = ap.parse_args()

    tz = ZoneInfo(args.tz)
    cutoff = datetime.now(timezone.utc) - timedelta(days=int(args.years * 365.25))
    screenshot_dir = Path(args.screenshot_dir or f"{args.out_prefix}_screenshots")

    raw_records = load_records(args.input)
    seen, rows = set(), []
    dropped_old = dropped_no_match = dropped_bad_date = 0
    screenshot_attempted = screenshot_ok = 0

    for r in raw_records:
        key = (r.get("account"), r.get("message_id"))
        if key in seen:
            continue
        seen.add(key)

        subject = r.get("subject") or ""
        body = r.get("body_text") or ""
        body_html = r.get("body_html") or ""
        matched = find_matches(subject + "\n" + body)
        if not matched:
            dropped_no_match += 1
            continue

        dt_utc = parse_date_any(r.get("date_raw"))
        if dt_utc is None:
            dropped_bad_date += 1
            continue
        if dt_utc < cutoff:
            dropped_old += 1
            continue

        dt_local = dt_utc.astimezone(tz)
        hour = dt_local.hour
        is_flagged = hour >= args.flag_start_hour or hour < args.flag_end_hour
        date_local_str = dt_local.strftime("%Y-%m-%d %I:%M %p %Z")

        screenshot_path = None
        want_shot = args.screenshots == "all" or (args.screenshots == "flagged" and is_flagged)
        if want_shot:
            screenshot_attempted += 1
            screenshot_dir.mkdir(parents=True, exist_ok=True)
            candidate_path = screenshot_dir / f"{len(rows)}.png"
            screenshot_path = render_email_screenshot(
                subject,
                r.get("sender_name", ""),
                r.get("sender_email", ""),
                date_local_str,
                body_html,
                body,
                candidate_path,
            )
            if screenshot_path:
                screenshot_ok += 1

        rows.append({
            "date_local": dt_local,
            "date_local_str": date_local_str,
            "day_of_week": dt_local.strftime("%A"),
            "account": r.get("account", ""),
            "folder": r.get("folder", ""),
            "sender_name": r.get("sender_name", ""),
            "sender_email": r.get("sender_email", ""),
            "company_guess": guess_company(r.get("sender_email"), r.get("sender_name")),
            "subject": subject,
            "matched_keywords": ", ".join(matched),
            "potential_fdcpa_flag": "Yes" if is_flagged else "No",
            "excerpt": (body[:300] + "...") if len(body) > 300 else body,
            "source_ref": r.get("source_ref", ""),
            "screenshot_path": str(screenshot_path) if screenshot_path else "",
        })

    rows.sort(key=lambda x: x["date_local"], reverse=True)

    print(f"Loaded {len(raw_records)} raw records", file=sys.stderr)
    print(f"  Dropped (no keyword match on final check): {dropped_no_match}", file=sys.stderr)
    print(f"  Dropped (older than {args.years} years): {dropped_old}", file=sys.stderr)
    print(f"  Dropped (unparseable date): {dropped_bad_date}", file=sys.stderr)
    print(f"  Final matches: {len(rows)}", file=sys.stderr)
    print(f"  Potential FDCPA (off-hours) matches: "
          f"{sum(1 for r in rows if r['potential_fdcpa_flag'] == 'Yes')}", file=sys.stderr)
    if args.screenshots != "none":
        print(f"  Screenshots captured: {screenshot_ok}/{screenshot_attempted}", file=sys.stderr)

    write_outputs(rows, args.out_prefix)


if __name__ == "__main__":
    main()
